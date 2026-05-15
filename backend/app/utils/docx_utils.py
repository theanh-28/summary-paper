import logging
from docx import Document

logger = logging.getLogger(__name__)

def extract_text_from_docx(file_path: str) -> tuple[str, int]:
    """
    Trích xuất toàn bộ văn bản từ một file DOCX.
    
    Args:
        file_path (str): Đường dẫn tới file DOCX.
        
    Returns:
        tuple[str, int]: Nội dung text trích xuất được và số trang (mặc định là 1 vì DOCX không có trang cố định).
    """
    try:
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        
        # Có thể đọc text từ bảng (tables) nếu cần thiết
        # for table in doc.tables:
        #     for row in table.rows:
        #         for cell in row.cells:
        #             full_text.append(cell.text)
                    
        text = '\n'.join(full_text)
        return text, 1
    except Exception as e:
        logger.error(f"Lỗi khi đọc file DOCX {file_path}: {str(e)}")
        raise ValueError(f"Không thể trích xuất văn bản từ DOCX: {str(e)}")
